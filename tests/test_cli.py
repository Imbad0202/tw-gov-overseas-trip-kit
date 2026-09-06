"""使用者從初始化到產檔的流程，以及輸入錯誤不留下半套文件。"""
import json

import openpyxl
from docx import Document

from tripkit.cli import main


def _init(tmp_path):
    directory = tmp_path / "input with spaces"
    assert main(["init", "--directory", str(directory)]) == 0
    return directory


def _edit(path, edit):
    data = json.loads(path.read_text())
    edit(data)
    path.write_text(json.dumps(data, ensure_ascii=False))


def test_complete_user_journey(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    directory = _init(tmp_path)
    inputs = ["--trip", str(directory / "trip.json"), "--finance", str(directory / "finance.json")]
    assert main(["validate", *inputs, "--outputs", "report", "review", "handbook", "finance"]) == 0
    out = tmp_path / "result"
    assert main(["render", *inputs, "--outputs", "report", "review", "handbook", "finance", "--out", str(out)]) == 0
    assert main(["flights", "--input", str(directory / "flight-options.json"),
                 "--out", str(out / "flight_options.html")]) == 0
    assert len(list(out.iterdir())) == 5
    doc = Document(out / "report.docx")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "資料中心參訪" in text and "撰寫提示" not in text
    wb = openpyxl.load_workbook(out / "finance.xlsx")
    assert "合成試算" in wb.active["G3"].value
    wb.close()
    assert "轉機過夜" in (out / "flight_options.html").read_text()


def test_missing_body_requires_explicit_draft_mode(tmp_path, capsys):
    directory = _init(tmp_path)
    _edit(directory / "trip.json", lambda d: d.pop("report_content"))
    args = ["render", "--trip", str(directory / "trip.json"), "--outputs", "review", "report",
            "--out", str(tmp_path / "out")]
    assert main(args) == 2
    assert not (tmp_path / "out").exists()
    assert "report_content" in capsys.readouterr().err
    assert main([*args, "--allow-incomplete-report"]) == 0


def test_summary_is_not_required_for_handbook(tmp_path):
    directory = _init(tmp_path)
    _edit(directory / "trip.json", lambda d: d.pop("summary"))
    args = ["render", "--trip", str(directory / "trip.json"), "--out", str(tmp_path / "out")]
    assert main([*args, "--outputs", "handbook"]) == 0
    assert main([*args, "--outputs", "report", "--allow-incomplete-report"]) == 2


def test_malformed_json_and_date_errors_are_actionable(tmp_path, capsys):
    bad = tmp_path / "bad.json"
    bad.write_text('{"agency":')
    assert main(["validate", "--trip", str(bad)]) == 2
    assert "JSON 格式錯誤" in capsys.readouterr().err
    directory = _init(tmp_path)
    _edit(directory / "trip.json", lambda d: d["trip"].update(start_date="2027-13-99"))
    assert main(["validate", "--trip", str(directory / "trip.json")]) == 2
    assert "$.trip.start_date" in capsys.readouterr().err


def test_overwrite_protection(tmp_path):
    directory = _init(tmp_path)
    path = directory / "trip.json"
    path.write_text("user work")
    assert main(["init", "--directory", str(directory)]) == 2
    assert path.read_text() == "user work"
    assert main(["init", "--directory", str(directory), "--force"]) == 0
    out = tmp_path / "out"
    out.mkdir()
    report = out / "review_table.docx"
    report.write_text("edited by user")
    args = ["render", "--trip", str(path), "--out", str(out)]
    assert main(args) == 2
    assert report.read_text() == "edited by user"
    assert not (out / "pre_trip.html").exists()
    assert main([*args, "--force"]) == 0


def test_pair_mismatch(tmp_path, capsys):
    directory = _init(tmp_path)
    _edit(directory / "finance.json", lambda d: d["agency"].update(full_name="另一示範機關"))
    assert main(["validate", "--trip", str(directory / "trip.json"),
                 "--finance", str(directory / "finance.json")]) == 2
    assert "不一致" in capsys.readouterr().err


def test_flight_input_shape_and_clock(tmp_path):
    path = tmp_path / "bad-flight.json"
    for data in ({"candidates": [None]}, {"candidates": [], "first_duty_local": "25:00"}):
        path.write_text(json.dumps(data))
        assert main(["flights", "--input", str(path)]) == 2
