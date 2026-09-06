"""經費規劃表 xlsx。日支明細 + B 類 manual + 總計（R6 進位）+ 動態簽章列。"""
import openpyxl
from openpyxl.styles import Alignment, Font
from calc.per_diem import compute_trip_per_diem


def render_finance_xlsx(data: dict, out_path: str) -> None:
    """產生經費規劃表 xlsx。

    Args:
        data: 包含 agency、per_diem_inputs、signatures 的 dict。
        out_path: 輸出 xlsx 路徑。
    """
    agency = data["agency"]
    inputs = data["per_diem_inputs"]
    result = compute_trip_per_diem(
        inputs["segments"],
        inputs.get("manual_items"),
        approved_days=inputs.get("approved_days"),
        approved_start_date=inputs.get("approved_start_date"),
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "經費規劃表"

    # 標題列
    ws["A1"] = f"{agency['full_name']} 出國經費規劃表"

    # 表頭
    ws.append(["日期", "留宿地", "日支基準(US$)", "當日應領(US$)", "計算條件", "備註／核定依據", "數額來源"])

    # 逐日日支明細
    for seg in result["per_segment"]:
        conditions = [{"none": "自理膳宿", "board_only": "供膳不供宿",
                       "lodging_only": "供宿不供膳", "board_and_lodging": "供膳宿"}
                      [seg.get("host_provided", "none")]]
        if seg.get("is_return_day"):
            conditions.append("返國日30%")
        if seg.get("day_index_same_place", 1) > 30:
            conditions.append(f"同地第{seg['day_index_same_place']}日（30/90日模型）")
        if seg.get("reduction_exempt"):
            conditions.append("遞減豁免")
        if seg.get("approved_extension"):
            conditions.append("已核准延返")
        if seg.get("cash_allowance_usd"):
            conditions.append(f"現金津貼US$ {seg['cash_allowance_usd']}")
        if seg.get("meals_not_provided"):
            meals = {"breakfast": "早", "lunch": "午", "dinner": "晚"}
            conditions.append("未供餐：" + "、".join(meals[m] for m in seg["meals_not_provided"]))
        ws.append([
            seg.get("date", ""),
            seg.get("lodging_city", ""),
            seg["per_diem_base"],
            seg["amount_usd"],
            "；".join(conditions),
            seg.get("note", ""),
            seg.get("rate_source", ""),
        ])

    # B 類 manual_items
    for m in inputs.get("manual_items", []):
        ws.append([m["label"], "(人工填入)", "", m["amount_usd"], "", m.get("note", "")])

    # 總計列（R6 尾數進位整數）
    ws.append(["總計(US$，尾數進位整數)", "", "", result["grand_total_usd"]])
    total_row = ws.max_row
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
    for cell in ws[total_row]:
        cell.font = Font(bold=True)

    # 動態簽章列（從 data['signatures'] 陣列生成，不寫死角色）
    ws.append([])
    ws.append([s["role"] for s in data.get("signatures", [])])
    ws.append([s.get("name", "") for s in data.get("signatures", [])])
    ws.append(["美元規劃試算；修改輸入後須重新產檔。總計使用未逐日進位金額，最後才進位整數。"])
    ws.merge_cells("A1:G1")
    ws["A1"].font = Font(size=16, bold=True)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=7)
    ws.freeze_panes = "C3"
    for column, width in {"A": 24, "B": 20, "C": 18, "D": 20, "E": 38, "F": 45, "G": 45}.items():
        ws.column_dimensions[column].width = width
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                # 使用者文字即使以 = 開頭也不應變成 Excel 公式。
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.column in (3, 4) and isinstance(cell.value, (int, float)):
                cell.number_format = "0.00####"
    for cell in ws[2]:
        cell.font = Font(bold=True)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:2"

    wb.save(out_path)
